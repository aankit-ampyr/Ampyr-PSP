"""
Dump key sheet structures from the financial Excel model.
Extracts cell addresses, labels, values, and formulas for mapping.

Usage: python _dump_financial_cells.py
Output: Structured report of the 5 key sheets needed for financial_config.py
"""

import openpyxl
from openpyxl.utils import get_column_letter

WB_PATH = r"Financial Model/Off-Grid Solution v8.xlsm"

# Sheets and row ranges to inspect
SHEET_SPECS = {
    "Solar&BESS Inputs": {"rows": None, "cols": (1, 20)},      # Full sheet — input parameters
    "Equity":            {"rows": (1, 200), "cols": (1, 30)},   # FCFF + XIRR rows
    "FS":                {"rows": None, "cols": (1, 30)},       # Revenue, OPEX, EBITDA
    "D&T":               {"rows": None, "cols": (1, 30)},       # Depreciation & Tax
    "Construction":      {"rows": None, "cols": (1, 30)},       # CAPEX phasing
}


def get_label(ws, row, col):
    """Try to find a label for this cell by looking at columns to the left."""
    for offset in range(1, min(col, 4)):
        candidate = ws.cell(row=row, column=col - offset).value
        if candidate and isinstance(candidate, str) and not str(candidate).startswith("="):
            return candidate.strip()
    return ""


def classify_cell(value):
    """Classify cell content type."""
    if value is None:
        return "empty"
    if isinstance(value, str) and value.startswith("="):
        return "formula"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, str):
        return "text"
    return type(value).__name__


def dump_sheet(ws, sheet_name, row_range=None, col_range=(1, 20)):
    """Dump non-empty cells from a sheet with labels and types."""
    min_row = row_range[0] if row_range else ws.min_row
    max_row = row_range[1] if row_range else min(ws.max_row, 500)
    min_col = col_range[0]
    max_col = col_range[1]

    print(f"\n{'=' * 100}")
    print(f"SHEET: {sheet_name}")
    print(f"Range: rows {min_row}-{max_row}, cols {get_column_letter(min_col)}-{get_column_letter(max_col)}")
    print(f"{'=' * 100}")
    print(f"{'Row':<6} {'Cell':<8} {'Type':<10} {'Label':<40} {'Value/Formula'}")
    print(f"{'-' * 6} {'-' * 8} {'-' * 10} {'-' * 40} {'-' * 50}")

    row_has_content = False
    last_row = None

    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            cell = ws.cell(row=row, column=col)
            value = cell.value

            if value is None:
                continue

            cell_type = classify_cell(value)

            # Skip pure whitespace text
            if cell_type == "text" and not str(value).strip():
                continue

            cell_ref = f"{get_column_letter(col)}{row}"
            label = get_label(ws, row, col) if col > 1 else ""

            # Truncate long values for display
            display_val = str(value)
            if len(display_val) > 80:
                display_val = display_val[:77] + "..."

            # Add blank line between non-consecutive rows
            if last_row is not None and row > last_row + 1:
                print()

            print(f"{row:<6} {cell_ref:<8} {cell_type:<10} {label:<40} {display_val}")
            last_row = row


def main():
    print(f"Loading workbook: {WB_PATH}")
    print("(formulas visible, not computed values)\n")

    wb = openpyxl.load_workbook(WB_PATH, read_only=False, data_only=False, keep_vba=True)

    print(f"Sheets in workbook: {wb.sheetnames}\n")

    for sheet_name, spec in SHEET_SPECS.items():
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            dump_sheet(
                ws,
                sheet_name,
                row_range=spec.get("rows"),
                col_range=spec.get("cols", (1, 20)),
            )
        else:
            # Try fuzzy match
            matches = [s for s in wb.sheetnames if sheet_name.lower() in s.lower()]
            if matches:
                print(f"\n*** Sheet '{sheet_name}' not found. Did you mean: {matches}?")
            else:
                print(f"\n*** Sheet '{sheet_name}' not found. Available: {wb.sheetnames}")

    # Also dump sheet names with dimensions for reference
    print(f"\n{'=' * 100}")
    print("ALL SHEETS SUMMARY")
    print(f"{'=' * 100}")
    for name in wb.sheetnames:
        ws = wb[name]
        print(f"  {name:<30} rows: {ws.min_row}-{ws.max_row}, cols: {ws.min_column}-{ws.max_column}")

    wb.close()


if __name__ == "__main__":
    main()
