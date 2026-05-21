import openpyxl
import json

wb_path = r"C:\repos\Ampyr-PSP\Financial Model\Off-Grid Solution v8.xlsm"

# Load with formulas visible
wb = openpyxl.load_workbook(wb_path, read_only=False, data_only=False, keep_vba=True)

print("=" * 80)
print("SHEET INVENTORY")
print("=" * 80)
for i, name in enumerate(wb.sheetnames):
    ws = wb[name]
    print(f"\n--- Sheet {i+1}: '{name}' ---")
    print(f"  Dimensions: {ws.dimensions}")
    print(f"  Min/Max Row: {ws.min_row}-{ws.max_row}")
    print(f"  Min/Max Col: {ws.min_column}-{ws.max_column}")
    
    # Count formulas vs values
    formula_count = 0
    value_count = 0
    empty_count = 0
    for row in ws.iter_rows(min_row=ws.min_row, max_row=min(ws.max_row, 1000), 
                             min_col=ws.min_column, max_col=min(ws.max_column, 100)):
        for cell in row:
            if cell.value is None:
                empty_count += 1
            elif isinstance(cell.value, str) and cell.value.startswith('='):
                formula_count += 1
            else:
                value_count += 1
    print(f"  Formulas (sampled): {formula_count}")
    print(f"  Values (sampled): {value_count}")
    print(f"  Empty (sampled): {empty_count}")

print("\n" + "=" * 80)
print("DEFINED NAMES (Named Ranges)")
print("=" * 80)
for dn in wb.defined_names.definedName:
    if not dn.name.startswith('_') and not dn.name.startswith('\'):
        destinations = list(dn.destinations)
        if destinations:
            for sheet_title, cell_range in destinations:
                print(f"  {dn.name} -> '{sheet_title}'!{cell_range}")
        else:
            print(f"  {dn.name} -> {dn.attr_text}")

print("\n" + "=" * 80)
print("DATA VALIDATIONS (Dropdowns / Input Controls)")
print("=" * 80)
for name in wb.sheetnames:
    ws = wb[name]
    if ws.data_validations and ws.data_validations.dataValidation:
        for dv in ws.data_validations.dataValidation:
            cells = str(dv.sqref) if dv.sqref else "unknown"
            formula1 = dv.formula1 if dv.formula1 else ""
            dv_type = dv.type if dv.type else ""
            print(f"  Sheet '{name}' | Cells: {cells} | Type: {dv_type} | Values: {formula1}")

wb.close()
