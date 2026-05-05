"""Dump every cell on the Project IRR computation path so we can compare
against Python, line by line. NOT for equity/debt repayment logic."""

import openpyxl
from openpyxl.utils import get_column_letter

WB = "Financial Model/Off-Grid Solution v8.xlsm"

SHEETS_AND_RANGES = [
    # Solar+BESS Project IRR — Equity sheet rows 145-180 (FCFF chain + XIRR)
    ("Equity", (145, 180), (1, 12)),
    # FS sheet — revenue/opex/capex feed rows
    ("FS", (1, 100), (1, 12)),
    # D&T sheet — tax rows used in FCFF
    ("D&T", (180, 270), (1, 12)),
    # Construction sheet — CAPEX phasing
    ("Construction", (75, 130), (1, 12)),
    # Consol Cash Flows — combined Project IRR
    ("Consol Cash Flows", (1, 25), (1, 15)),
    # Cash Flows-Gas — gas FCFF
    ("Cash Flows-Gas", (1, 90), (1, 12)),
    # Capex & Funding-Gas — feeds into Cash Flows-Gas
    ("Capex & Funding-Gas", (1, 100), (1, 12)),
    # SETUP — global switches
    ("SETUP", (1, 60), (1, 10)),
    # Overall Inputs
    ("Overall Inputs", (1, 80), (1, 10)),
    # Curves and D&T — tax rates
    ("Curves and D&T", (100, 120), (1, 8)),
]


def dump(ws, name, row_range, col_range):
    print(f"\n{'='*100}\nSHEET: {name}  rows {row_range[0]}-{row_range[1]}, cols {get_column_letter(col_range[0])}-{get_column_letter(col_range[1])}\n{'='*100}")
    last_row = None
    for row in range(row_range[0], min(row_range[1] + 1, ws.max_row + 1)):
        for col in range(col_range[0], min(col_range[1] + 1, ws.max_column + 1)):
            v = ws.cell(row=row, column=col).value
            if v is None:
                continue
            s = str(v)
            if len(s) > 110:
                s = s[:107] + "..."
            if last_row is not None and row > last_row + 1:
                print()
            print(f"  {get_column_letter(col)}{row}: {s}")
            last_row = row


def main():
    wb = openpyxl.load_workbook(WB, data_only=False, keep_vba=True)
    print(f"Sheets: {wb.sheetnames}")
    for name, rr, cr in SHEETS_AND_RANGES:
        if name not in wb.sheetnames:
            print(f"\n*** Sheet '{name}' missing")
            continue
        dump(wb[name], name, rr, cr)
    wb.close()


if __name__ == "__main__":
    main()
